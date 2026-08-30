"""Safe resume parsing for TXT, PDF, and DOCX uploads."""

from __future__ import annotations

from dataclasses import dataclass
from hashlib import sha256
from io import BytesIO
from pathlib import Path
import re


MAX_FILE_BYTES = 8 * 1024 * 1024
SUPPORTED_EXTENSIONS = {".txt", ".pdf", ".docx"}


class ResumeParseError(ValueError):
    """Raised when an uploaded resume cannot be converted to usable text."""


@dataclass(frozen=True)
class ParsedResume:
    filename: str
    text: str
    file_type: str
    page_count: int
    word_count: int
    resume_id: str


def redact_personal_data(text: str) -> str:
    """Remove common contact details before text is sent to external AI services."""

    redacted = re.sub(
        r"[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}",
        "[EMAIL REDACTED]",
        text,
    )
    redacted = re.sub(
        r"(?<!\w)(?:\+?\d[\d ()-]{7,}\d)(?!\w)",
        "[PHONE REDACTED]",
        redacted,
    )
    redacted = re.sub(
        r"https?://(?:www\.)?(?:linkedin\.com|github\.com)/\S+",
        "[PROFILE URL REDACTED]",
        redacted,
        flags=re.IGNORECASE,
    )
    return redacted


def _clean_text(text: str) -> str:
    text = text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in text.splitlines()]
    cleaned = "\n".join(lines)
    cleaned = re.sub(r"\n{3,}", "\n\n", cleaned).strip()
    return cleaned


def _parse_txt(data: bytes) -> tuple[str, int]:
    for encoding in ("utf-8-sig", "utf-8", "cp1252"):
        try:
            return data.decode(encoding), 1
        except UnicodeDecodeError:
            continue
    raise ResumeParseError("The text file encoding could not be detected.")


def _parse_pdf(data: bytes) -> tuple[str, int]:
    try:
        from pypdf import PdfReader
    except ImportError as error:
        raise ResumeParseError("PDF support is unavailable. Install pypdf.") from error

    try:
        reader = PdfReader(BytesIO(data))
        pages = [(page.extract_text() or "") for page in reader.pages]
    except Exception as error:
        raise ResumeParseError(f"The PDF could not be read: {error}") from error
    return "\n\n".join(pages), len(reader.pages)


def _parse_docx(data: bytes) -> tuple[str, int]:
    try:
        from docx import Document
    except ImportError as error:
        raise ResumeParseError("DOCX support is unavailable. Install python-docx.") from error

    try:
        document = Document(BytesIO(data))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                paragraphs.append(" | ".join(cell.text for cell in row.cells))
    except Exception as error:
        raise ResumeParseError(f"The DOCX file could not be read: {error}") from error
    return "\n".join(paragraphs), 1


def parse_resume(data: bytes, filename: str) -> ParsedResume:
    """Extract and normalize resume text without writing the upload to disk."""

    if not data:
        raise ResumeParseError("The uploaded file is empty.")
    if len(data) > MAX_FILE_BYTES:
        raise ResumeParseError("The resume is larger than the 8 MB upload limit.")

    extension = Path(filename).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise ResumeParseError(f"Unsupported file type. Upload one of: {supported}.")

    if extension == ".txt":
        raw_text, page_count = _parse_txt(data)
    elif extension == ".pdf":
        raw_text, page_count = _parse_pdf(data)
    else:
        raw_text, page_count = _parse_docx(data)

    text = _clean_text(raw_text)
    word_count = len(re.findall(r"\b[\w+#.-]+\b", text))
    if word_count < 25:
        raise ResumeParseError(
            "Very little text was extracted. If this is a scanned PDF, run OCR first.",
        )

    return ParsedResume(
        filename=Path(filename).name,
        text=text,
        file_type=extension.removeprefix(".").upper(),
        page_count=page_count,
        word_count=word_count,
        resume_id=sha256(data).hexdigest()[:20],
    )
